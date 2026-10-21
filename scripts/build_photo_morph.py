"""Build 22 individual still backgrounds, separate Morph drone, native charts; zero video."""
from pathlib import Path
import csv,json,subprocess,copy,math,random,shutil
from PIL import Image,ImageOps,ImageDraw,ImageFilter
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
import imageio_ffmpeg
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1];AS=ROOT/'presentations/static_morph_assets';AS.mkdir(exist_ok=True);OUT=ROOT/'presentations/AGROVECTOR_PHOTO_MORPH.pptx'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
SOURCE=ROOT/'presentations/AGROVECTOR_CINEMATIC.pptx'
def cover(im,i):
 im=ImageOps.fit(im.convert('RGB'),(1920,1080),centering=(.5,.45 if i==2 else .5));return im

def prepare():
 if Path('/tmp/agrovector-wheat-eye-level.jpg').exists():shutil.copyfile('/tmp/agrovector-wheat-eye-level.jpg',AS/'eye_level_original.jpg')
 ff=imageio_ffmpeg.get_ffmpeg_exe();sources=['/tmp/agrovector-field-horizon.mp4','/tmp/agrovector-9289796.mp4','/tmp/agrovector-5308472.mp4'];library=[]
 lib=AS/'source_library.json'
 if not lib.exists():lib=Path('/tmp/agrovector-photo-library.json')
 if lib.exists():
  raw=json.loads(lib.read_text());library=raw if isinstance(raw,list) else raw.get('photos',raw.get('assets',[]))
 ledger=[]
 for i in range(1,23):
  dst=AS/f'background_{i:02d}.jpg'
  if i==2:
   im=Image.open(AS/'eye_level_original.jpg');source={'author':'Barney Goodman','source_url':'https://unsplash.com/photos/detailed-golden-wheat-ears-under-a-bright-blue-sky-sMndTj2ZfsQ','license_url':'https://unsplash.com/license'}
  elif i>=3 and library:
   item=library[(i-3)%len(library)];path=item.get('filename',item.get('file',item.get('path','')));path=Path(path)
   if not path.is_absolute():path=(AS/path) if (AS/path).exists() else Path('/tmp/agrovector-photo-library')/path
   im=Image.open(path);source=item
   # Separate frames/compositions, not an animated background.
   if i-3>=len(library):
    w,h=im.size;cycle=(i-3)//len(library);inset=.06+.055*cycle;im=im.crop((int(w*inset),int(h*inset*.8),int(w*(1-inset*.8)),int(h*(1-inset*.8))))
  else:
   index=0 if i==1 else (i%3);time=2 if i==1 else (i%6)*.7
   temp=AS/f'raw_{i:02d}.jpg'
   if not temp.exists():subprocess.run([ff,'-y','-loglevel','error','-ss',str(time),'-i',sources[index],'-frames:v','1',str(temp)],check=True)
   im=Image.open(temp)
   source={'source_url':['https://www.pexels.com/video/aerial-footage-of-wheat-field-9130276/','https://www.pexels.com/video/aerial-footage-of-a-wheat-field-9289796/','https://www.pexels.com/video/close-up-view-of-a-wheat-5308472/'][index],'license_url':'https://www.pexels.com/license/','type':'Selected STILL frame, no video embedded'}
  cover(im,i).save(dst,quality=94);ledger.append({'slide':i,'background':str(dst.relative_to(ROOT)),'source':source,'motion':'subtle native photo drift' if i==1 else 'NONE: static photo'})
 # Only mist, no second baked-in drone. It shares the native drone's motion path.
 rng=random.Random(21);haze=Image.new('RGBA',(1200,1100));d=ImageDraw.Draw(haze)
 for j in range(2600):
  t=rng.random();x=600+rng.gauss(0,100+180*t);y=500+580*t;alpha=int((1-t)*38+8);d.ellipse((x,y,x+3,y+6),fill=(230,244,246,alpha))
 haze=haze.filter(ImageFilter.GaussianBlur(1.6));haze.save(AS/'spray_only.png')
 (AS/'photo_ledger.json').write_text(json.dumps(ledger,ensure_ascii=False,indent=2)+'\n');return ledger

def chart(slide,kind,categories,series,title,name):
 from build_cinematic_presentation import text,panel
 panel(slide,7.35,2.85,5.35,3.75,alpha=95000)
 text(slide,7.6,2.96,4.8,.35,title,12,True,'D6ED89')
 data=CategoryChartData();data.categories=categories
 for label,values in series:data.add_series(label,values)
 shape=slide.shapes.add_chart(kind,Inches(7.6),Inches(3.45),Inches(4.85),Inches(2.8),data);shape.name=name;ch=shape.chart;ch.has_legend=False
 if kind==XL_CHART_TYPE.DOUGHNUT:
  ch.plots[0].hole_size=65
  for i,point in enumerate(ch.series[0].points):point.format.fill.solid();point.format.fill.fore_color.rgb=RGBColor.from_string(['D6ED89','E7AE65','83BBB1'][i%3])
 else:
  for ax in (ch.category_axis,ch.value_axis):ax.tick_labels.font.size=Pt(10);ax.tick_labels.font.color.rgb=RGBColor(255,255,255)
  for j,s in enumerate(ch.series):s.format.fill.solid();s.format.fill.fore_color.rgb=RGBColor.from_string(['D6ED89','E7AE65','83BBB1'][j%3])
 return shape

def build():
 ledger=prepare();p=Presentation(str(SOURCE));base=Presentation(str(ROOT/'presentations/AGROVECTOR_22_SLIDES.pptx'));plan=[]
 for i,s in enumerate(p.slides,1):
  # Remove the movie AND relationships, preventing orphan embedded MP4s.
  for sh in list(s.shapes):
   videos=sh._element.xpath('.//a:videoFile')
   if videos:
    rids=set()
    for x in sh._element.iter():
     for key,val in x.attrib.items():
      if key.startswith('{'+R+'}'):rids.add(val)
    s.shapes._spTree.remove(sh._element)
    for rid in rids:
     if rid in s.part.rels:s.part.drop_rel(rid)
  bg=s.shapes.add_picture(str(AS/f'background_{i:02d}.jpg'),Inches(-.08 if i==1 else 0),Inches(-.05 if i==1 else 0),width=Inches(13.49 if i==1 else 13.333333),height=Inches(7.6 if i==1 else 7.5));bg.name='!!AV_background';tree=s.shapes._spTree;tree.remove(bg._element);tree.insert(2,bg._element)
  # Static backgrounds everywhere except a tiny native motion on the first photo.
  x,y,w=(7.2,1.35,5.5) if i==1 else ((9.0,2.1,3.05) if i==2 else (11.42,.79,1.12))
  drone=s.shapes.add_picture(str(ROOT/'presentations/cinematic_assets/drone_render.png'),Inches(x),Inches(y),width=Inches(w));drone.name='!!AV_drone'
  haze=None
  if i==2:
   haze=s.shapes.add_picture(str(AS/'spray_only.png'),Inches(x),Inches(y),width=Inches(w));haze.name='!!AV_spray';tree=s.shapes._spTree;tree.remove(haze._element);tree.insert(list(tree).index(drone._element),haze._element)
  for sh in s.shapes:
   if sh.has_text_frame and 'Визуализация концепции + сток Pexels.' in sh.text:
    for par in sh.text_frame.paragraphs:
     for run in par.runs:run.text=run.text.replace('Визуализация концепции + сток Pexels.','Иллюстративные фото: Unsplash/Pexels.')
  if i in (19,20):
   cats=['CAPEX','Запуск','Ликвидность'] if i==19 else ['Учредитель','Инвестор']
   vals=[7,.5,2] if i==19 else [2.5,7]
   chart(s,XL_CHART_TYPE.DOUGHNUT,cats,[('План, млн руб.',vals)],'ПРОЕКТНЫЕ СУММЫ, НЕ ДОЛИ В КАПИТАЛЕ' if i==20 else 'СТАРТОВАЯ ПОТРЕБНОСТЬ: 9,5 МЛН РУБ.',f'AV_DONUT_{i}')
   from build_cinematic_presentation import text,panel
   for j,(label,value) in enumerate(zip(cats,vals)):
    xx=7.53+j*(1.65 if i==19 else 2.45);tint=['D6ED89','E7AE65','83BBB1'][j]
    panel(s,xx,6.4,.07,.1,tint,100000);text(s,xx+.12,6.3,1.5 if i==19 else 2.3,.35,label+' '+str(value).replace('.',','),10,False,'FFFFFF')
  if i==11:
   rows=list(csv.DictReader((ROOT/'data/financial_model_5y.csv').open()))[1:]
   vals=[float(r['treatment_ha_passes'])/130+float(r['monitoring_ha_surveys'])/500 for r in rows]
   chart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,[r['year'] for r in rows],[('Полевые дни',vals)],'ПЛАНОВЫЕ ПОЛЕВЫЕ ДНИ; МОЩНОСТЬ НЕ ПОДТВЕРЖДЕНА','AV_COLUMNS_11')
  s.notes_slide.notes_text_frame.text=base.slides[i-1].notes_slide.notes_text_frame.text+'\n\nФОТО + MORPH, РЕДАКЦИЯ ВИЗУАЛА. Фон — статичное иллюстративное фото, не видео и не съемка работ AgroVector. Источник: '+json.dumps(ledger[i-1]['source'],ensure_ascii=False)+'\nДрон — отдельный PNG-объект !!AV_drone с одинаковым именем на слайдах, для нативного Morph. Первое фото едва покачивается; на втором движутся только дрон и связанное распыление. Диаграммы отражают расчеты/проектные суммы, не подтвержденные продажи или доли капитала.'
  plan.append({'slide':i,'background_id':bg.shape_id,'drone_id':drone.shape_id,'spray_id':haze.shape_id if haze else None,'chart_ids':[sh.shape_id for sh in s.shapes if sh.has_chart]})
 p.save(str(OUT));(AS/'shape_plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2)+'\n');print(OUT)
if __name__=='__main__':build()
