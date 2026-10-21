"""True Morph matching + native PNG motion paths + per-point chart entrances."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import json,copy
from lxml import etree as E
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];FILE=ROOT/'presentations/AGROVECTOR_PHOTO_MORPH.pptx';P='http://schemas.openxmlformats.org/presentationml/2006/main';A='http://schemas.openxmlformats.org/drawingml/2006/main';MC='http://schemas.openxmlformats.org/markup-compatibility/2006';M='http://schemas.microsoft.com/office/powerpoint/2015/09/main';NS={'p':P,'a':A}
def sub(parent,tag,**attrs):return E.SubElement(parent,'{'+P+'}'+tag,{k:str(v) for k,v in attrs.items()})
class Timeline:
 def __init__(self):
  self.count=1;self.timing=E.Element('{'+P+'}timing');par=sub(sub(self.timing,'tnLst'),'par');root=sub(par,'cTn',id=self.uid(),dur='indefinite',restart='never',nodeType='tmRoot');self.rootchildren=sub(root,'childTnLst');seq=sub(self.rootchildren,'seq',concurrent=1,nextAc='seek');main=sub(seq,'cTn',id=self.uid(),dur='indefinite',nodeType='mainSeq');self.children=sub(main,'childTnLst');self.builds=[]
  for name,evt in [('prevCondLst','onPrev'),('nextCondLst','onNext')]:sub(sub(sub(seq,name),'cond',evt=evt,delay=0),'tgtEl').append(E.Element('{'+P+'}sldTgt'))
 def uid(self):n=self.count;self.count+=1;return n
 def effect(self,delay=0,preset=None):
  par=sub(self.children,'par');attrs={'id':self.uid(),'nodeType':'withEffect','fill':'hold'}
  if preset is not None:attrs.update(presetID=preset,presetClass='entr',grpId=0)
  ct=sub(par,'cTn',**attrs);sub(sub(ct,'stCondLst'),'cond',delay=delay);return sub(ct,'childTnLst')
 def target(self,beh,spid,point=None):
  t=sub(sub(beh,'tgtEl'),'spTgt',spid=spid)
  if point is not None:
   series,category=point;g=sub(t,'graphicEl');E.SubElement(g,'{'+A+'}chart',bldStep='ptInSeries',seriesIdx=str(series),categoryIdx=str(category))
 def visibility(self,children,spid,point):
  st=sub(children,'set');beh=sub(st,'cBhvr');ct=sub(beh,'cTn',id=self.uid(),dur=1,fill='hold');sub(sub(ct,'stCondLst'),'cond',delay=0);self.target(beh,spid,point);sub(sub(beh,'attrNameLst'),'attrName').text='style.visibility';sub(sub(st,'to'),'strVal',val='visible')
 def chartpoint(self,spid,series,category,filter,delay):
  c=self.effect(delay,21 if filter=='wheel(1)' else 22);point=(series,category);self.visibility(c,spid,point);ae=sub(c,'animEffect',transition='in',filter=filter);beh=sub(ae,'cBhvr');sub(beh,'cTn',id=self.uid(),dur=650 if filter=='wheel(1)' else 700,fill='hold');self.target(beh,spid,point)
 def chartbuild(self,spid):self.builds.append(spid)
 def motion(self,spid,path,duration):
  par=sub(self.rootchildren,'par');ct=sub(par,'cTn',id=self.uid(),fill='hold');sub(sub(ct,'stCondLst'),'cond',delay=0);c=sub(ct,'childTnLst');motion=sub(c,'animMotion',origin='layout',path=path,pathEditMode='relative');beh=sub(motion,'cBhvr',additive='base');ct=sub(beh,'cTn',id=self.uid(),dur=duration,autoRev=1,repeatCount='indefinite',accel=40000,decel=40000,fill='hold');sub(sub(ct,'stCondLst'),'cond',delay=0);self.target(beh,spid);attrs=sub(beh,'attrNameLst');sub(attrs,'attrName').text='ppt_x';sub(attrs,'attrName').text='ppt_y'
 def finish(self):
  if not len(self.children):self.rootchildren.remove(self.children.getparent().getparent())
  if self.builds:
   lst=sub(self.timing,'bldLst')
   for spid in self.builds:E.SubElement(sub(sub(lst,'bldGraphic',spid=spid,grpId=0),'bldSub'),'{'+A+'}bldChart',bld='seriesEl',animBg='0')
  return self.timing

def build():
 prs=Presentation(str(FILE));plan=json.loads((ROOT/'presentations/static_morph_assets/shape_plan.json').read_text());report=[]
 with ZipFile(FILE) as z:parts={n:z.read(n) for n in z.namelist()}
 for i,(sl,rec) in enumerate(zip(prs.slides,plan),1):
  name=f'ppt/slides/slide{i}.xml';tree=E.fromstring(parts[name])
  for child in list(tree):
   if E.QName(child).localname in ('timing','transition','AlternateContent'):tree.remove(child)
  # Same exact drone PNG/name across slides; Morph actually transports the object.
  ac=E.Element('{'+MC+'}AlternateContent',nsmap={'mc':MC});choice=E.SubElement(ac,'{'+MC+'}Choice',Requires='p159',nsmap={'p159':M});tr=E.SubElement(choice,'{'+P+'}transition',spd='slow',advClick='1');E.SubElement(tr,'{'+M+'}morph',option='byObject');fallback=E.SubElement(ac,'{'+MC+'}Fallback');sub(sub(fallback,'transition',spd='med',advClick=1),'fade')
  idx=next((j for j,x in enumerate(tree) if E.QName(x).localname=='extLst'),len(tree));tree.insert(idx,ac)
  tl=Timeline();animations=[]
  if i==1:
   tl.motion(rec['background_id'],'M 0 0 L 0 -0.004 E',5000);tl.motion(rec['drone_id'],'M 0 0 L 0 -0.012 E',2200);animations=['photo subtle drift','separate drone hover']
  if i==2:
   tl.motion(rec['drone_id'],'M 0 0 L -0.12 0 E',4000)
   if rec['spray_id']:tl.motion(rec['spray_id'],'M 0 0 L -0.12 0 E',4000)
   animations=['drone flight + synchronized spray; photo static']
  count=0
  for sh in sl.shapes:
   if not sh.has_chart:continue
   chart=sh.chart;kind=chart.chart_type;donut='DOUGHNUT' in str(kind)
   # Grow every vertical column from baseline. Mixed-sign horizontal sensitivity
   # bars wipe outwards from zero, according to sign, not uniformly from one edge.
   for si,series in enumerate(chart.series):
    for ci,value in enumerate(series.values):
     filt='wheel(1)' if donut else ('wipe(left)' if value<0 else 'wipe(right)') if i==18 else 'wipe(up)'
     delay=ci*700 if donut else ci*170+si*95
     tl.chartpoint(sh.shape_id,si,ci,filt,delay);count+=1
   tl.chartbuild(sh.shape_id);animations.append('sector Wheel entrances' if donut else 'individual point Wipe entrances')
  if animations:tree.insert(idx+1,tl.finish())
  parts[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True);report.append({'slide':i,'morph':True,'animations':animations,'animated_chart_points':count,'background_motion':i==1})
 with ZipFile(FILE,'w',ZIP_DEFLATED) as z:
  for name,data in parts.items():z.writestr(name,data)
 (ROOT/'presentations/static_morph_assets/animation_plan.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':build()
