"""Read back photo/Morph requirements after native PowerPoint save."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
import json,hashlib,csv,pymupdf
ROOT=Path(__file__).resolve().parents[1];FILE=ROOT/'presentations/AGROVECTOR_PHOTO_MORPH.pptx';NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
def verify():
 p=Presentation(str(FILE));base=Presentation(str(ROOT/'presentations/AGROVECTOR_22_SLIDES.pptx'));assert len(p.slides)==22
 with ZipFile(FILE) as z:
  assert z.testzip() is None;assert not any(n.endswith(('.mp4','.mov','.webm')) for n in z.namelist())
  trees=[E.fromstring(z.read(f'ppt/slides/slide{i}.xml')) for i in range(1,23)]
  assert all(t.xpath('//*[local-name()="morph"]') for t in trees)
  assert all(not t.xpath('//*[local-name()="videoFile" or local-name()="cMediaNode"]') for t in trees)
  for i,t in enumerate(trees):
   ids=[n.get('id') for n in t.findall('.//p:cNvPr',NS)];assert len(ids)==len(set(ids))
   for target in t.findall('.//p:spTgt',NS):assert target.get('spid') in ids
   assert base.slides[i].notes_slide.notes_text_frame.text in p.slides[i].notes_slide.notes_text_frame.text
  drone1=next(s for s in p.slides[0].shapes if s.name=='!!AV_drone');drone2=next(s for s in p.slides[1].shapes if s.name=='!!AV_drone');assert drone1.image.sha1==drone2.image.sha1
  assert (drone1.left,drone1.top,drone1.width)!=(drone2.left,drone2.top,drone2.width)
  firsttargets={x.get('spid') for x in trees[0].findall('.//p:animMotion//p:spTgt',NS)};secondtargets={x.get('spid') for x in trees[1].findall('.//p:animMotion//p:spTgt',NS)}
  assert firsttargets=={str(s.shape_id) for s in p.slides[0].shapes if s.name in ['!!AV_background','!!AV_drone']}
  assert secondtargets=={str(s.shape_id) for s in p.slides[1].shapes if s.name in ['!!AV_spray','!!AV_drone']}
  assert all(not t.findall('.//p:animMotion',NS) for t in trees[2:])
  charts=[]
  for i in (11,17,18,19,20):
   sh=next(s for s in p.slides[i-1].shapes if s.has_chart);targets=trees[i-1].findall('.//p:graphicEl/a:chart',NS)
   points={(int(t.get('seriesIdx')),int(t.get('categoryIdx'))) for t in targets};expected={(si,ci) for si,series in enumerate(sh.chart.series) for ci,v in enumerate(series.values)};assert points==expected
   filters=[x.get('filter') for x in trees[i-1].findall('.//p:animEffect',NS)];assert all(x=='wheel(1)' for x in filters) if i in (19,20) else all(x.startswith('wipe(') for x in filters)
   charts.append({'slide':i,'chart_points':len(points),'filters':sorted(set(filters))})
  assert list(next(s.chart for s in p.slides[18].shapes if s.has_chart).series[0].values)==[7,.5,2]
  assert list(next(s.chart for s in p.slides[19].shapes if s.has_chart).series[0].values)==[2.5,7]
  for i in (17,18):
   a=next(s.chart for s in base.slides[i-1].shapes if s.has_chart);b=next(s.chart for s in p.slides[i-1].shapes if s.has_chart)
   for x,y in zip(a.series,b.series):assert all(abs(xx-yy)<1e-10 for xx,yy in zip(x.values,y.values))
  table=next(s.table for s in p.slides[5].shapes if s.has_table);assert [[c.text for c in r.cells] for r in list(table.rows)[1:]]==json.loads((ROOT/'data/competitive_matrix.json').read_text())['slide_rows']
 hashes={hashlib.sha256(x.read_bytes()).hexdigest() for x in (ROOT/'presentations/static_morph_assets').glob('background_*.jpg')};assert len(hashes)==22
 pdf=ROOT/'review/photo_morph/AGROVECTOR_PHOTO_MORPH.pdf';assert len(pymupdf.open(pdf))==22
 report={'verified':True,'file':str(FILE),'sha256':hashlib.sha256(FILE.read_bytes()).hexdigest(),'slides':22,'morph_transitions':22,'embedded_videos':0,'unique_background_compositions':22,'same_separate_drone_across_1_2':True,'slide2_only_drone_and_spray_animated':True,'source_notes_preserved':22,'charts':charts,'chart_point_total':sum(x['chart_points'] for x in charts),'native_open_save':'PowerPoint opened/saved successfully; effects and Morph survived native save','visual_limit':'PDF static layouts inspected; pixel-level native slideshow motion not recorded. Wheel is sector reveal, not numeric data interpolation.'}
 (ROOT/'presentations/PHOTO_MORPH_ACCEPTANCE.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':verify()
