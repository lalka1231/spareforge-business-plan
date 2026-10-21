"""AgroVector editorial redesign: light analytical pages + dark milestones."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'presentations/AGROVECTOR_22_SLIDES.pptx'
OUT=ROOT/'presentations/AGROVECTOR_22_SLIDES_DESIGN.pptx'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
NS={'a':A}
DARK={1,5,11,17,18,22}
def build():
 p=Presentation(str(SRC))
 for i,s in enumerate(p.slides,1):
  light=i not in DARK
  bg='F5F3EC' if light else '101F1B'; s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(bg)
  for sh in s.shapes:
   for node in sh._element.xpath('.//a:srgbClr'):
    old=node.get('val'); ancestor=node.getparent(); text=False
    while ancestor is not None:
     if E.QName(ancestor).localname in ('defRPr','rPr','endParaRPr'):text=True;break
     ancestor=ancestor.getparent()
    if light and text:
     tc=next((x for x in node.iterancestors() if E.QName(x).localname=='tc'),None)
     darkcell=tc is not None and any(x.get('val') in ('164C38','092A20') for x in tc.findall('a:tcPr/a:solidFill/a:srgbClr',NS))
     colors={'FFFFFF':'182F28' if not darkcell else 'FFFFFF','C8F26B':'386349' if not darkcell else 'D7EF8B','AFC5B8':'65766D','092A20':'182F28','164C38':'274D3B'}
    else:colors={'092A20':'101F1B','164C38':'224C3B','C8F26B':'D7EF8B','EFF4EF':'F4F2E9','AFC5B8':'B4C4B8'}
    node.set('val',colors.get(old,old))
   # Sharper editorial panels instead of repeated rounded app cards.
   for geo in sh._element.xpath('.//a:prstGeom'):
    if geo.get('prst')=='roundRect':geo.set('prst','rect')
   if sh.has_text_frame:
    for latin in sh._element.xpath('.//a:latin'):latin.set('typeface','Aptos')
  # Stable framing geometry that Morph can match across slides.
  tag=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(13.12),Inches(.4),Inches(.07),Inches(6.5)); tag.name='!!AV_editorial_edge';tag.fill.solid();tag.fill.fore_color.rgb=RGBColor.from_string('54735A' if light else 'D7EF8B');tag.line.fill.background()
 # Original vector illustration: decorative, not a map or equipment specification.
 s=p.slides[0]
 for k in range(9):
  line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(7.35+k*.54),Inches(5.74),Inches(8.35+k*.43),Inches(4.66));line.line.color.rgb=RGBColor.from_string('3C5D46');line.line.width=Pt(1.2)
 # Stylized quadcopter silhouette floats over field lines.
 cx,cy=9.95,5.05
 for dx,dy in [(-.8,-.22),(.8,-.22),(-.8,.22),(.8,.22)]:
  line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(cx),Inches(cy),Inches(cx+dx),Inches(cy+dy));line.line.color.rgb=RGBColor.from_string('D7EF8B');line.line.width=Pt(3)
  rotor=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(cx+dx-.26),Inches(cy+dy-.09),Inches(.52),Inches(.18));rotor.fill.background();rotor.line.color.rgb=RGBColor.from_string('D7EF8B');rotor.line.width=Pt(1.8)
 body=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(cx-.22),Inches(cy-.14),Inches(.44),Inches(.28));body.fill.solid();body.fill.fore_color.rgb=RGBColor.from_string('D7EF8B');body.line.fill.background()
 # Keep all original data/notes/tables/charts. Only decorative elements are new.
 p.save(str(OUT));check=Presentation(str(OUT));assert len(check.slides)==22
 for x,y in zip(Presentation(str(SRC)).slides,check.slides):
  assert [sh.text for sh in x.shapes if sh.has_text_frame and sh.text]==[sh.text for sh in y.shapes if sh.has_text_frame and sh.text]
  assert x.notes_slide.notes_text_frame.text==y.notes_slide.notes_text_frame.text
 print(OUT)
if __name__=='__main__':build()
