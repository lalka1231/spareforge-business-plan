"""Revise the existing PHOTO_MORPH deck against the supplied teaching reference."""
from pathlib import Path
import json,csv
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from build_cinematic_presentation import text,panel,table,color
from build_presentation import sections_for_notes
ROOT=Path(__file__).resolve().parents[1];FILE=ROOT/'presentations/AGROVECTOR_PHOTO_MORPH.pptx'
def load(n):return json.loads((ROOT/'data'/n).read_text())
def rows(n):return list(csv.DictReader((ROOT/'data'/n).open()))
def m(v):return f'{v/1e6:.3f}'.replace('.',',')
def pb(v):return 'не достигнута за 5 лет' if v is None else f'{v:.2f} года'.replace('.',',')
def card(s,x,y,w,h,title,body,size=15):
 title=title.replace('Уникальное ценностное предложение','Уникальная ценность').replace('Труднокопируемое преимущество','Преимущество')
 panel(s,x,y,w,h,alpha=97000);text(s,x+.13,y+.10,w-.26,.42,title,13,True,'D6ED89');sh=text(s,x+.13,y+.59,w-.26,h-.64,body,size)
 for p in sh.text_frame.paragraphs:p.space_after=Pt(5)
 return sh
def line(s,x1,y1,x2,y2):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=color('D6ED89');sh.line.width=Pt(1.5)
def clear(s,n,title,subtitle):
 for sh in list(s.shapes):
  if sh.name not in ('!!AV_background','!!AV_drone','!!AV_spray'):s.shapes._spTree.remove(sh._element)
 # A quiet, opaque analytical plane preserves the field scene without reducing legibility.
 panel(s,.35,.28,12.65,6.69,alpha=96000)
 text(s,.62,.43,11.8,.35,f'АГРОВЕКТОР / {n:02d} / {title.upper()}',13,True,'D6ED89')
 text(s,.62,1.00,12.0,.60,subtitle,27,True)
 text(s,.62,7.13,11.7,.22,'Проектные гипотезы и расчеты; не подтвержденный бизнес. Референс: Практика 08.10.26. / Редакция 21.10.2026',9,False,'CCDACF')
 text(s,12.12,7.11,.7,.24,f'{n:02d}/22',10,True,'D6ED89')
def chart(s,kind,cats,series,x,y,w,h):
 d=CategoryChartData();d.categories=cats
 for name,vals in series:d.add_series(name,vals)
 sh=s.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),d);ch=sh.chart;ch.has_legend=False
 for ax in (ch.category_axis,ch.value_axis):ax.tick_labels.font.size=Pt(11);ax.tick_labels.font.color.rgb=color('FFFFFF')
 if kind==XL_CHART_TYPE.BAR_CLUSTERED:
  plot=ch.plots[0];plot.has_data_labels=True;plot.data_labels.number_format='0.00';plot.data_labels.font.size=Pt(11);plot.data_labels.font.color.rgb=color('FFFFFF')
 for j,ser in enumerate(ch.series):
  ser.format.fill.solid();ser.format.fill.fore_color.rgb=color(['D6ED89','E5B26A','A7CAB9'][j%3]);ser.invert_if_negative=False
 return sh

def build():
 p=Presentation(str(FILE));ref=load('reference_revision.json');f=load('finance_summary.json');b=f['base'];annual=f['yearly_scenarios']['base'];y=annual[1];sections=sections_for_notes();market=load('market_revision.json');prices=load('price_benchmarks.json');pay=load('payroll_revision.json');tasks=rows('project_plan.csv')
 s=p.slides[3];clear(s,4,'Lean Canvas','Бизнес-модель: девять проверяемых гипотез')
 # Exactly the reference layout: five columns; solution/metrics and advantage/channels split.
 col=2.40;xx=[.62+j*2.44 for j in range(5)];top=1.91
 content={r[0].split('.')[0]:r for r in ref['lean']}
 concise={'2':'Агроокна; неоднородность посевов; нехватка карт.\n\nАльтернативы: наземная техника, свой БВС, подрядчик.','4':'Обследование → карта → внесение → отчет.\nПлатный пилот и приемка.','3':'От карты проблемы до согласованной операции и проверяемого отчета.\n\nГипотеза ценности, не обещание урожайности.','9':'Пока НЕ подтверждено.\nБудущий актив: карты и повторные договоры.','1':'КФХ и малые/средние хозяйства Татарстана.\nПокупатель — директор; приемка — агроном.\n\nРанние: готовые оплатить конкретный пилот.','8':'Оплаченные га; вклад; повторы; отчет ≤48 ч.\nЦель:10 клиентов.','5':'Прямые продажи; агрономы; рекомендации; демонстрации.','7':f'Штат + взносы: {m(pay["annual_employer_cost_rub"])} млн ₽.\nПрочий OPEX; CAPEX и резерв отдельно.','6':'Внесение: 6 000 × 1 300 = 7,8 млн ₽.\nМониторинг: 4 000 × 250 = 1,0 млн ₽. Всего 8,8 млн ₽.'}
 for key,c,r,h in [('2',0,top,3.26),('4',1,top,1.55),('8',1,top+1.68,1.58),('3',2,top,3.26),('9',3,top,1.55),('5',3,top+1.68,1.58),('1',4,top,3.26)]:card(s,xx[c],r,col,h,content[key][0],concise[key],14)
 for key,x in [('7',.62),('6',6.76)]:card(s,x,5.33,5.96,1.36,content[key][0],concise[key],15)
 s=p.slides[4];clear(s,5,'Рынок: TAM / SAM / SOM','TAM — вся Россия; SAM — доступный кластер; SOM — наш план')
 for j,k in enumerate(['tam','sam','som']):
  z=market[k];card(s,.62+j*4.09,1.97,3.92,3.74,z['label'],z['visible'],20)
 text(s,.66,5.98,11.8,.67,market['limitation'],15,False,'D6ED89')
 s=p.slides[5];clear(s,6,'Ценовые ориентиры аналогичных компаний','Публичные ориентировочные цены — не сопоставимые КП')
 table(s,['Компания / регион','Операция / цена','Что входит и ограничения'],[[z['company'],z['price'],z['scope']] for z in prices['rows']],y=1.98,h=3.68,widths=[2.6,3.35,6.15],size=15)
 text(s,.67,5.94,11.9,.75,'AgroVector: план 1 300 ₽/га-проход и 250 ₽/га-обследование; препараты клиента.\nДоли конкурентов публично не раскрыты. Полное нейтральное сравнение, URL и даты — в заметках.',16,False,'D6ED89')
 s=p.slides[6];clear(s,7,'Развернутый SWOT / TOWS','Факторы среды → четыре стратегии действий')
 sw=ref['swot_factors'];st=ref['swot_strategies']
 data=[['Внутренние / внешние','O — возможности\nЛокальные сложные поля; платные пилоты с агрономами; подготовка кадров','T — угрозы\nОграничения работ; погода и отказы; цена и задержки платежей'],['S — проектные сильные стороны\nКонтур «карта–операция–отчет»; резервный БВС; короткие решения','SO: S1 + O1/O2\nПлатный пилот с паспортом поля и критериями приемки.\nДиректор + агроном; до закупки.','ST: S2/S3 + T1/T2\nДопуск, резерв, чек-лист; no-go без оснований.\nОпер. руководитель; каждый выезд.'],['W — слабые стороны\nНет заказов и сезона; одна бригада; парк/кадры не подтверждены','WO: W1/W3 + O2/O3\nКП, договорный портфель, предложения кандидатам и наставник.\nДиректор; до финансирования.','WT: W2 + T2/T3\nЛимит месячной загрузки, аванс и переносы; отказ от перегруза.\nДиректор; перед договором.']]
 table(s,data[0],data[1:],y=1.96,h=4.36,widths=[3.6,4.25,4.25],size=15)
 text(s,.66,6.48,11.8,.28,'Гипотезы раннего проекта. Полный перечень S1–T3, метрики и ответственные — в заметках и письменном плане.',12,False,'D6ED89')
 s=p.slides[7];clear(s,8,'Пять сил Портера','Сила давления, объяснение и направление работ')
 table(s,['Параметр','Уровень','Описание','Направление работ'],ref['porter'],y=1.96,h=4.63,widths=[2.25,1.4,3.65,4.8],size=13)
 s=p.slides[9];clear(s,10,'Стейкхолдеры / модель Менделоу','Власть × интерес определяют стратегию взаимодействия')
 for x,y0,w,h,t,body in [(.9,2.0,5.72,2.06,'Высокая власть / низкий интерес','Регуляторы, страховщик: поддерживать удовлетворенность.\nЮрист/директор — до допуска и при изменениях.'),(6.78,2.,5.72,2.06,'Высокая власть / высокий интерес','Учредитель, инвестор, ключевые клиенты: тесное управление.\nДиректор — еженедельно и перед решениями.'),(.9,4.2,5.72,2.06,'Низкая власть / низкий интерес','Не вовлеченные хозяйства вне маршрута, широкая аудитория: наблюдать.\nДиректор — ежемесячно.'),(6.78,4.2,5.72,2.06,'Низкая власть / высокий интерес','Операторы, агроном, кандидаты, соседи/пчеловоды: информировать.\nОпер. руководитель — до выезда; уведомления по закону.')]:card(s,x,y0,w,h,t,body,17)
 text(s,.4,3.1,.45,1.5,'В\nЛ\nА\nС\nТ\nЬ',12,True,'D6ED89');text(s,5.65,6.48,3,.27,'ИНТЕРЕС →',13,True,'D6ED89')
 s=p.slides[11];clear(s,12,'Команда и стоимость труда','Три штатные единицы: gross и взносы раздельно')
 table(s,['Роль','FTE','Gross / мес.','Gross / год','Взносы / год','Работодатель / год'],pay['slide_rows'],y=1.95,h=2.44,widths=[3.05,.6,1.8,2.1,2.1,2.45],size=15)
 card(s,.65,4.62,5.91,1.75,'РЕГИОНАЛЬНЫЙ ОРИЕНТИР',pay['benchmark_visible'],17)
 card(s,6.74,4.62,5.91,1.75,'РАСЧЕТ И ОГРАНИЧЕНИЯ','30% общие взносы + 0,2% травматизм (допущение).\nНДФЛ внутри gross. Агроном, юрист и бухгалтер — отдельно в OPEX.',17)
 text(s,.66,6.55,11.9,.25,'Оклады — план найма, не подтвержденные предложения кандидатам. Региональная средняя не равна зарплате пилота.',11,False,'D6ED89')
 s=p.slides[12];clear(s,13,'WBS / дерево работ','От запуска сервиса к проверяемым результатам')
 panel(s,4.86,1.88,3.6,.52);text(s,4.98,1.97,3.35,.3,'0. Запуск AgroVector',17,True,'D6ED89')
 groups=[('1. Проверка',['A','B','C']),('2. Ресурсы',['D','E','F','G']),('3. Готовность',['H','I','J','M']),('4. Сезон / аудит',['K','L'])];short={'A':'Правовой маршрут','B':'Интервью и заявки','C':'ООО и договоры','D':'Решение / транши','E':'Приемка техники','F':'Подготовка команды','G':'Документы / no-go','H':'SOP / карты / CRM','I':'Стендовая приемка','J':'Договоры / маршруты','M':'Предсезонный допуск','K':'Пилот + сезон','L':'Аудит / продление'}
 line(s,6.66,2.4,6.66,2.62);line(s,2.08,2.62,11.25,2.62)
 for j,(title,ids) in enumerate(groups):
  x=.65+j*3.075;line(s,x+1.43,2.62,x+1.43,2.82);card(s,x,2.82,2.89,.68,title,'',14)
  for k,tid in enumerate(ids):
   yy=3.65+k*.62;panel(s,x+.13,yy,2.64,.5,alpha=98000);text(s,x+.23,yy+.10,2.43,.32,tid+'  '+short[tid],14)
 text(s,.66,6.48,11.8,.30,'Все 13 пакетов из project_plan.csv. Это дерево декомпозиции, не календарь: зависимости — на слайде 15.',12,False,'D6ED89')
 s=p.slides[15];clear(s,16,'Матрица ответственности / RACI','Кто выполняет, утверждает, консультирует и получает информацию')
 roles=['Директор','Инвестор','Юрист','Бухгалтер','Агроном','Опер. рук.','Оператор 2','Учебный партнер','Клиент'];alias={'Второй оператор':'Оператор 2','Продажи':'Директор','Операционный руководитель':'Опер. рук.','Клиенты':'Клиент'};rr=[]
 for task in tasks:
  cells={r:[] for r in roles}
  for code in ['R','A','C','I']:
   for raw in task[code].split(';'):
    targets=['Опер. рук.','Оператор 2'] if raw=='Команда' else [alias.get(raw,raw)]
    for role in targets:
     if role in cells and code not in cells[role]:cells[role].append(code)
  rr.append([task['id']+' '+short[task['id']]]+['/'.join(cells[r]) or '—' for r in roles])
 tab=table(s,['Работа']+roles,rr,y=1.94,h=4.29,widths=[3.18]+[.99]*9,size=10)
 for ri,row in enumerate(tab.rows):
  row.height=Inches(.44 if ri==0 else .285)
  for cell in row.cells:
   cell.margin_top=cell.margin_bottom=Inches(.015)
   for para in cell.text_frame.paragraphs:para.space_before=Pt(0);para.space_after=Pt(0)
 for i in range(1,len(rr)+1):
  for j in range(1,10):
   cell=tab.cell(i,j);v=cell.text
   if 'A' in v:tint='E5B26A'
   elif 'R' in v:tint='D6ED89'
   elif 'C' in v:tint='A7CAB9'
   else:tint='FFFFFF'
   for para in cell.text_frame.paragraphs:para.font.color.rgb=color(tint);para.font.bold='A' in v or 'R' in v;para.font.size=Pt(12)
 text(s,.67,6.41,12,.32,'R — выполняет   A — утверждает   C — консультирует   I — информируется   — — не участвует',13,False,'D6ED89')
 text(s,.67,6.78,11.9,.22,'Опер. рук. = первый оператор. Оператор 2 = второй сотрудник; в I/K выполнение и утверждение разделены.',11,False,'FFFFFF')
 s=p.slides[16];clear(s,17,'Финансовая модель','Повышенная стоимость труда снижает финансовый результат')
 card(s,.65,1.97,4.43,3.91,'2027 / МЛН РУБ.',f'Выручка: {m(y["revenue_rub"])}\nПеременные: {m(y["variable_cost_rub"])}\nФикс. OPEX: {m(y["fixed_opex_rub"])}\nEBITDA: {m(y["ebitda_rub"])}\nFCF: {m(y["free_cash_flow_rub"])}',21)
 chart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,[str(r['year']) for r in annual[1:]],[('Выручка',[r['revenue_rub']/1e6 for r in annual[1:]]),('EBITDA',[r['ebitda_rub']/1e6 for r in annual[1:]]),('FCF',[r['free_cash_flow_rub']/1e6 for r in annual[1:]])],5.34,2.05,7.1,3.85)
 text(s,5.6,6.05,6.6,.4,'Светлый — выручка; охра — EBITDA; мятный — FCF',13,False,'D6ED89')
 s=p.slides[17];clear(s,18,'NPV и чувствительность','Новая зарплатная база — новый инвестиционный вывод')
 card(s,.65,1.97,4.1,4.4,'СТАВКА 24% / 5 ЛЕТ',f'NPV: {m(b["npv_rub"])} млн ₽\n\nПростая окупаемость:\n{pb(b["simple_payback_years"])}\n\nДисконтированная:\n{pb(b["discounted_payback_years"])}',19)
 sens=f['sensitivity'];labels=['База','Тариф −10%','Объем −25%','Переменные +20%','Задержка 1 год','Внесение 900 ₽']
 chart(s,XL_CHART_TYPE.BAR_CLUSTERED,labels,[('NPV, млн ₽',[r['npv_rub']/1e6 for r in sens])],4.96,2.02,7.55,4.25)
 s=p.slides[20];clear(s,21,'Ключевые показатели','Финансы и мощность: расчет не равен подтвержденной готовности')
 table(s,['Показатель','2027','2031'],[['Выручка, млн ₽',m(annual[1]['revenue_rub']),m(annual[5]['revenue_rub'])],['EBITDA, млн ₽',m(annual[1]['ebitda_rub']),m(annual[5]['ebitda_rub'])],['FCF, млн ₽',m(annual[1]['free_cash_flow_rub']),m(annual[5]['free_cash_flow_rub'])],['Внесение / мониторинг','6000 / 4000','13000 / 10000'],['Бригады / FTE','1 / 3','1 / 3']],y=2.0,h=3.8,widths=[6.1,3,3],size=18)
 text(s,.67,6.02,11.9,.65,'Поздние объемы ресурсно не подтверждены. FCF 2031 включает возврат резерва 2 млн ₽ только в конце горизонта.',16,False,'D6ED89')
 # Every note is regenerated from the same current written section, not inherited stale text.
 for i,s in enumerate(p.slides,1):
  for sh in s.shapes:
   if sh.has_text_frame and 'Редакция 20.10.2026' in sh.text:
    for para in sh.text_frame.paragraphs:
     for run in para.runs:run.text=run.text.replace('Редакция 20.10.2026','Редакция 21.10.2026')
  s.notes_slide.notes_text_frame.text=sections[str(i)]+'\n\nРЕДАКЦИЯ ПО РЕФЕРЕНСУ 21.10.2026. Фото иллюстративные, не реальные работы AgroVector. Морф переносит отдельный PNG !!AV_drone. Источники изображений: presentations/static_morph_assets/photo_ledger.json. Исходный учебный референс: Практика 08.10.26.pptx.'
 plan=[]
 for i,s in enumerate(p.slides,1):
  sh={z.name:z for z in s.shapes};plan.append({'slide':i,'background_id':sh['!!AV_background'].shape_id,'drone_id':sh['!!AV_drone'].shape_id,'spray_id':sh['!!AV_spray'].shape_id if '!!AV_spray' in sh else None,'chart_ids':[z.shape_id for z in s.shapes if z.has_chart]})
 p.save(str(FILE));(ROOT/'presentations/static_morph_assets/shape_plan.json').write_text(json.dumps(plan,indent=2)+'\n');print(FILE)
if __name__=='__main__':build()
