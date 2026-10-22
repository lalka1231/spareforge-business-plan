"""Read-only acceptance of the after-brainstorm deliverables."""
from pathlib import Path
from zipfile import ZipFile
import hashlib,json,re,subprocess
from lxml import etree as E
from docx import Document
from pptx import Presentation
import pymupdf
from build_finance import model,metrics
ROOT=Path(__file__).resolve().parents[1]
PARENT='428ed6b1ea89876a6a910299f843efaa8507062e'
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}

def verify():
    changed=subprocess.check_output(['git','diff',PARENT,'--name-only'],cwd=ROOT,text=True).splitlines()
    prior=set(subprocess.check_output(['git','ls-tree','-r','--name-only',PARENT],cwd=ROOT,text=True).splitlines())
    assert not prior.intersection(changed),'Historical files changed: '+str(prior.intersection(changed))
    current=json.loads((ROOT/'data/finance_corrected_brainstorm.json').read_text())
    baseline=json.loads((ROOT/'data/finance_summary.json').read_text())
    assert current['base']==metrics(model())==baseline['base']
    assert current['yearly_scenarios']['base']==baseline['yearly_scenarios']['base']
    cash=current['cash_runway'];reserve=model()[0]['reserve_funding_rub'];burn=model()[1]['fixed_opex_rub']/12
    assert burn==cash['no_revenue_fixed_burn_per_month_rub']==490600
    first=next(i for i in range(1,120) if reserve-i*burn<0)
    assert first==cash['no_revenue_first_negative_month']==5
    assert reserve-4*burn>=0 and reserve-5*burn<0
    md=(ROOT/'docs/FINAL_BUSINESS_PLAN_brainstorm.md').read_text()
    sections=re.findall(r'^## (\d+)\. ([\s\S]*?)(?=^## \d+\.|^## Sources|\Z)',md,re.M)
    assert [int(x[0]) for x in sections]==list(range(1,23))
    assert '{fmt(' not in md and 'Не достигнута года' not in md
    doc=Document(str(ROOT/'documents/AGROVECTOR_FULL_BUSINESS_PLAN_brainstorm.docx'))
    dt='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    numbered=[p.text for p in doc.paragraphs if p.style and p.style.name=='Heading 1' and re.match(r'^\d+\.',p.text)]
    assert len(numbered)==22 and len(doc.inline_shapes)==2 and '{fmt(' not in dt
    ppt=Presentation(str(ROOT/'presentations/AGROVECTOR_AFTER_brainstorm.pptx'));assert len(ppt.slides)==22
    for (number,section),slide in zip(sections,ppt.slides):
        note=slide.notes_slide.notes_text_frame.text
        assert section in note,(number,'Notes are not synchronized to revised report')
        assert '{fmt(' not in note
    for marker in ('490 600','21 760','AgroVector_After_brainstorm'):
        assert marker in md and marker in dt,marker
    annual=current['yearly_scenarios']['base'][1:]
    finance_charts=[s.chart for s in ppt.slides[16].shapes if s.has_chart]
    assert len(finance_charts)==1
    for ser,key in zip(finance_charts[0].series,('revenue_rub','ebitda_rub','free_cash_flow_rub')):
        assert all(abs(a-b[key]/1e6)<1e-8 for a,b in zip(ser.values,annual))
    stress_charts=[s.chart for s in ppt.slides[17].shapes if s.has_chart]
    assert len(stress_charts)==1
    assert all(abs(a-b['npv_rub']/1e6)<1e-8 for a,b in zip(stress_charts[0].series[0].values,baseline['sensitivity']))
    chart_count=0
    with ZipFile(ROOT/'presentations/AGROVECTOR_PHOTO_MORPH.pptx') as old,ZipFile(ROOT/'presentations/AGROVECTOR_AFTER_brainstorm.pptx') as new:
        assert new.testzip() is None
        for i in range(1,23):
            path=f'ppt/slides/slide{i}.xml';a=E.fromstring(old.read(path));b=E.fromstring(new.read(path))
            if i<=2:assert old.read(path)==new.read(path)
            for tag in ('transition','timing'):
                assert [E.tostring(x) for x in a.xpath('//p:'+tag,namespaces=NS)]==[E.tostring(x) for x in b.xpath('//p:'+tag,namespaces=NS)],(i,tag)
        for path in old.namelist():
            if path.startswith('ppt/media/'):assert old.read(path)==new.read(path),path
            if re.fullmatch(r'ppt/charts/chart\d+\.xml',path):
                a=E.fromstring(old.read(path));b=E.fromstring(new.read(path))
                assert [E.tostring(x,method='c14n') for x in a.xpath('//c:ser',namespaces=NS)]==[E.tostring(x,method='c14n') for x in b.xpath('//c:ser',namespaces=NS)],path
                chart_count+=1
    pdf_counts={}
    for relative,count in [('presentations/AGROVECTOR_AFTER_brainstorm.pdf',22),('documents/AGROVECTOR_FULL_BUSINESS_PLAN_brainstorm.pdf',None)]:
        pdf=pymupdf.open(ROOT/relative);assert len(pdf)>0
        if count:assert len(pdf)==count
        text='\n'.join(p.get_text() for p in pdf);assert '{fmt(' not in text
        assert '490600' in re.sub(r'\s+','',text)
        pdf_counts[relative]=len(pdf)
    assets=['docs/FINAL_BUSINESS_PLAN_brainstorm.md','data/finance_corrected_brainstorm.json','documents/AGROVECTOR_FULL_BUSINESS_PLAN_brainstorm.docx','presentations/AGROVECTOR_AFTER_brainstorm.pptx',*pdf_counts]
    for name in assets:assert (ROOT/name).stat().st_size<100*1024**2
    return {'verified':True,'parent_commit':PARENT,'historical_tracked_files_unchanged':True,'sections':22,'slides':22,'notes_synchronized':22,'timing_transition_preserved':22,'first_two_slides_identical':True,'chart_series_unchanged':chart_count,'word_tables':len(doc.tables),'word_diagrams':len(doc.inline_shapes),'pdf_pages':pdf_counts,'npv_rub':current['base']['npv_rub'],'fixed_monthly_rub':burn,'first_negative_no_revenue_month':first,'hashes':{name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for name in assets},'limits':['Static LibreOffice render; no new native PowerPoint slideshow playback test','External legal, tax, market, staffing, supplier, operational and financing gates remain open']}

if __name__=='__main__':print(json.dumps(verify(),ensure_ascii=False,indent=2))
