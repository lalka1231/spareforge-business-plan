"""Create a complete Word dossier from the checked business plan and CSVs.
Requires python-docx, markdown-it-py, Pillow. No Word rendering is implied.
"""
from pathlib import Path
import csv,json,re,math,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from markdown_it import MarkdownIt
from PIL import Image,ImageDraw,ImageFont
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'documents';ASSETS=OUT/'assets'
GREEN='164C38';LIGHT='EAF2EC'
def rows(name):return list(csv.DictReader((ROOT/'data'/name).open(encoding='utf-8')))
def num(v):
 try:
  x=float(v)
  if x.is_integer():return f'{x:,.0f}'.replace(',',' ')
  return f'{x:,.2f}'.replace(',',' ').replace('.',',')
 except (ValueError,TypeError):return str(v)
def field(p,code):
 r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),code);r._r.addnext(fld)
def hyperlink(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),GREEN);pr.append(col);r.append(pr)
 tx=OxmlElement('w:t');tx.text=label;r.append(tx);h.append(r);p._p.append(h)
def graph_assets():
 ASSETS.mkdir(parents=True,exist_ok=True)
 fontfile=Path('/System/Library/Fonts/Supplemental/Arial.ttf')
 if not fontfile.exists():raise RuntimeError('Set a Cyrillic-capable font path for diagram generation')
 font=lambda n:ImageFont.truetype(str(fontfile),n)
 plan=rows('project_plan.csv');W=2300;im=Image.new('RGB',(W,1200),'white');d=ImageDraw.Draw(im)
 d.text((45,20),'График подготовки и первого сезона • календарные недели',font=font(36),fill='#'+GREEN)
 x0=940;scale=23;top=130
 for w in range(0,56,5):
  x=x0+w*scale;d.line((x,100,x,1120),fill='#d8dfd9',width=2);d.text((x-8,70),str(w),font=font(24),fill='#333333')
 for i,r in enumerate(plan):
  y=top+i*73;d.text((45,y),r['id']+'  '+r['task'],font=font(26),fill='#222222')
  x=x0+int(r['es'])*scale;end=x0+int(r['ef'])*scale
  d.rectangle((x,y,end,y+38),fill='#164C38' if r['critical']=='True' else '#8EB29A')
 d.text((45,1140),'Темный цвет — критический путь. Начало: 07.09.2026. Коммерческий сезон: 12.04–12.09.2027.',font=font(25),fill='#444444')
 im.save(ASSETS/'gantt.png')
 im=Image.new('RGB',(2300,950),'white');d=ImageDraw.Draw(im)
 positions={'A':(80,170),'B':(80,600),'C':(310,170),'D':(550,350),'E':(790,170),'F':(790,600),'G':(1030,170),'H':(1030,600),'I':(1270,350),'J':(1510,350),'M':(1750,350),'K':(1990,350),'L':(1990,650)};nw=180;nh=110
 d.text((40,25),'Сетевой график • зависимости «окончание–начало»',font=font(38),fill='#'+GREEN)
 for r in plan:
  tx,ty=positions[r['id']]
  for pred in filter(None,r['predecessors'].split(';')):
   x,y=positions[pred];a=(x+nw,y+nh/2);b=(tx,ty+nh/2)
   if pred=='K' and r['id']=='L':a=(x+nw/2,y+nh);b=(tx+nw/2,ty)
   d.line((a,b),fill='#75867c',width=5);ang=math.atan2(b[1]-a[1],b[0]-a[0]);tip=[b,(b[0]-20*math.cos(ang-.45),b[1]-20*math.sin(ang-.45)),(b[0]-20*math.cos(ang+.45),b[1]-20*math.sin(ang+.45))];d.polygon(tip,fill='#75867c')
 for r in plan:
  x,y=positions[r['id']];col='#164C38' if r['critical']=='True' else '#658d72';d.rounded_rectangle((x,y,x+nw,y+nh),radius=12,fill=col)
  d.text((x+20,y+12),r['id'],font=font(38),fill='white');d.text((x+20,y+62),r['duration_weeks']+' нед.',font=font(25),fill='white')
 d.text((45,850),'Критический путь: A → C → D → E → G → I → J → M → K → L. Продолжительность: 55 недель.',font=font(29),fill='#'+GREEN)
 im.save(ASSETS/'network.png')

class Builder:
 def __init__(self):
  self.doc=Document();self.caption=0;self.orientation='portrait';self.source_tables=[];self.section_no=0
  sec=self.doc.sections[0];self.page(sec,False);sec.different_first_page_header_footer=True
  styles=self.doc.styles
  for st in ['Normal','Body Text','List Bullet','List Number']:
   styles[st].font.name='Times New Roman';styles[st].font.size=Pt(12)
   styles[st].paragraph_format.line_spacing=1.15;styles[st].paragraph_format.space_after=Pt(6)
  for st,size in [('Title',30),('Subtitle',16),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
   styles[st].font.name='Times New Roman';styles[st].font.size=Pt(size);styles[st].font.color.rgb=RGBColor.from_string(GREEN)
  styles['Heading 1'].paragraph_format.page_break_before=True
  styles['Heading 1'].paragraph_format.keep_with_next=True
  p=sec.header.paragraphs[0];p.text='АГРОВЕКТОР  |  Бизнес-план 2027–2031';p.style='Caption'
  p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('АгроВектор  •  ');field(p,'PAGE')
  self.doc.core_properties.title='АгроВектор — полный бизнес-план с расчетами'
  self.doc.core_properties.author='Даниил Ваздаев'
  self.doc.core_properties.subject='Услуги агродронов, инвестиционная и операционная модель'
 def page(self,s,land):
  s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
  s.page_width=Cm(29.7 if land else 21);s.page_height=Cm(21 if land else 29.7)
  s.left_margin=Cm(2.2);s.right_margin=Cm(1.8);s.top_margin=Cm(2);s.bottom_margin=Cm(2)
 def orient(self,land):
  mode='landscape' if land else 'portrait'
  if mode!=self.orientation:
   sec=self.doc.add_section(WD_SECTION_START.NEW_PAGE);self.page(sec,land);sec.different_first_page_header_footer=False;self.orientation=mode
 def para(self,text='',style=None):
  p=self.doc.add_paragraph(style=style)
  for item in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
   r=p.add_run(item[2:-2] if item.startswith('**') else item.strip('`'))
   if item.startswith('**'):r.bold=True
  return p
 def table(self,headers,data,label=None):
  # Split wide tables into repeated-key panels; every source cell is retained.
  if len(headers)>6:
   for k in range(1,len(headers),5):
    idx=[0]+list(range(k,min(k+5,len(headers))))
    self.table([headers[x] for x in idx],[[r[x] for x in idx] for r in data],(label or 'Продолжение расчетной таблицы')+f' — блок {(k-1)//5+1}')
   return
  land=len(headers)>=5;self.orient(land)
  self.caption+=1;p=self.para(f'Таблица {self.caption}. '+(label or f'Данные раздела {self.section_no}'),'Caption');p.paragraph_format.keep_with_next=True
  t=self.doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
  width=(25.7 if land else 17)/len(headers)
  for c in t.columns:c.width=Cm(width)
  for i,x in enumerate(headers):t.rows[0].cells[i].text=x
  prop=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');prop.append(repeat)
  for row in data:
   cells=t.add_row().cells
   for i,text in enumerate(row):cells[i].text=text
  for ri,row in enumerate(t.rows):
   trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
   for c in row.cells:
    if ri==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),LIGHT);c._tc.get_or_add_tcPr().append(sh)
    for p in c.paragraphs:
     p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3)
     for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10 if len(headers)>3 else 11);r.bold=(ri==0)
  self.doc.add_paragraph();self.orient(False)
 def image(self,name,caption):
  self.orient(True);self.doc.add_picture(str(ASSETS/name),width=Cm(25));self.para(caption,'Caption');self.orient(False)
 def cover(self,titles):
  self.para('АГРОВЕКТОР','Title');self.para('БИЗНЕС-ПЛАН','Subtitle')
  self.para('Сервис применения сельскохозяйственных дронов для мониторинга посевов и точечного внесения удобрений, средств защиты растений и биопрепаратов','Subtitle')
  self.para('Полная версия с финансовой моделью, организационным планом и расчетными приложениями')
  self.para('Горизонт планирования: 2027–2031\nВерсия исходных данных: 05.09.2026\nАвтор: Даниил Ваздаев')
  self.para('Для рассмотрения агропартнером или инвестором. Проект ранней стадии; финансирование, контракты и разрешения не объявляются полученными.')
  self.doc.add_page_break();self.doc.add_heading('Содержание',level=1)
  for i,t in enumerate(titles,1):self.para(f'{i}. {t}')
  for t in ['Источники','Приложение А. Полная финансовая модель','Приложение Б. Помесячный денежный поток','Приложение В. Детальная единичная экономика','Приложение Г. Расчет всех стресс-сценариев','Приложение Д. Реестр параметров и допущений','Приложение Е. Инвестиционная проверка и открытые вопросы']:self.para(t)
  self.doc.add_heading('Правила чтения документа',level=1)
  self.para('[ФАКТ] — проверенный источник; [РАСЧЕТ] — результат вычислений; [ДОПУЩЕНИЕ] — проектная гипотеза. Маркер перед абзацем или таблицей относится ко всему блоку, если строка не имеет отдельного маркера. Ссылки [n] раскрыты в разделе «Источники».')
  self.para('Денежные значения приводятся в рублях РФ, если прямо не указаны миллионы. Га-проходы и га-обследования — разные оплаченные операции, не сумма уникальных гектаров. Числа округляются только для представления; финансовая модель пересчитывается без промежуточного округления.')
 def markdown(self,body):
  tokens=MarkdownIt().enable('table').parse(body);i=0;style=None
  while i<len(tokens):
   t=tokens[i]
   if t.type=='heading_open':
    text=tokens[i+1].content;level=int(t.tag[1]);match=re.match(r'(\d+)\.',text)
    if match:self.section_no=int(match[1])
    p=self.doc.add_heading(text,level=max(1,level-1));i+=3;continue
   if t.type=='table_open':
    j=i+1;rr=[];row=[]
    while tokens[j].type!='table_close':
     if tokens[j].type=='tr_open':row=[]
     elif tokens[j].type=='inline':row.append(self.inline_plain(tokens[j]))
     elif tokens[j].type=='tr_close':rr.append(row)
     j+=1
    self.source_tables.append(rr);self.table(rr[0],rr[1:]);i=j+1;continue
   if t.type=='fence':
    if t.info=='mermaid':
     self.image('gantt.png' if 'gantt' in t.content else 'network.png','[РАСЧЕТ] Графическая форма проектного плана; исходные зависимости и даты приведены в таблицах раздела.')
    else:self.para(t.content)
   elif t.type=='list_item_open':style='List Bullet'
   elif t.type=='list_item_close':style=None
   elif t.type=='inline':self.para(t.content,style)
   i+=1
 def inline_plain(self,t):
  if t.children:return ''.join(x.content for x in t.children if x.type in ['text','code_inline'])
  return t.content

def main():
 OUT.mkdir(exist_ok=True);graph_assets();b=Builder();titles=json.loads((ROOT/'data/section_titles.json').read_text());b.cover(titles)
 text=(ROOT/'docs/FINAL_BUSINESS_PLAN.md').read_text();body=text[text.index('## 1.'):].split('## Sources')[0];b.markdown(body)
 b.doc.add_heading('Источники',level=1)
 for line in (ROOT/'data/references_block.md').read_text().splitlines():
  match=re.match(r'\[(\d+)\] (https?://\S+) — (.*)',line)
  if match:
   sid,url,title=match.groups();p=b.para(f'[{sid}] {title}. Дата обращения: 05.09.2026. ');hyperlink(p,'Открыть источник',url)
   # Retain printable URLs with permitted line breaks after path delimiters.
   p=b.para(url.replace('/','/\u200b').replace('%','\u200b%'))
   for r in p.runs:r.font.size=Pt(8)
 f=json.loads((ROOT/'data/finance_summary.json').read_text());years=f['yearly_scenarios']['base'][1:]
 b.doc.add_heading('Приложение А. Полная финансовая модель',level=1)
 b.para('[РАСЧЕТ] Все показатели ниже прочитаны из той же финансовой модели. Стартовые потоки: оборудование −7 000 000; подготовка −500 000; резерв −2 000 000; итого FCF₀ = −9 500 000 руб. Номинальный t0 — 31.12.2026; это агрегированная оценка, не календарь траншей.')
 names={'treatment_ha_passes':'Внесение, га-проходы','monitoring_ha_surveys':'Мониторинг, га-обследования','treatment_price_ex_vat_rub':'Тариф внесения без НДС, руб.','monitoring_price_ex_vat_rub':'Тариф мониторинга без НДС, руб.','treatment_variable_unit_rub':'Переменные внесения, руб./га','monitoring_variable_unit_rub':'Переменные мониторинга, руб./га','treatment_revenue_rub':'Выручка внесения','monitoring_revenue_rub':'Выручка мониторинга','revenue_rub':'Выручка всего','variable_cost_rub':'Переменные затраты','contribution_rub':'Маржинальный доход','fixed_opex_rub':'Постоянный OPEX','ebitda_rub':'EBITDA','depreciation_rub':'Амортизация','ebit_rub':'EBIT','usn_tax_rub':'Налог УСН','accounting_profit_rub':'Расчетная чистая прибыль','vat_threshold_assumed_rub':'Порог освобождения от НДС','vat_collected_rub':'Предъявленный / собранный НДС','vat_remitted_rub':'Перечисленный НДС','customer_billings_including_vat_rub':'Счета клиентам с НДС','additional_capex_rub':'CAPEX замен','reserve_return_rub':'Возврат резерва','free_cash_flow_rub':'Свободный денежный поток','discount_factor':'Коэффициент дисконтирования','present_value_rub':'Приведенный поток','cumulative_cf_rub':'Накопленный FCF с t0','cumulative_discounted_cf_rub':'Накопленный PV с t0'}
 b.table(['Показатель']+[str(r['year']) for r in years],[[title]+[(f"{r[k]:.6f}".replace('.',',') if k=='discount_factor' else num(r[k])) for r in years] for k,title in names.items()],'Полная модель — рубли, кроме явно указанных единиц')
 b.para('[РАСЧЕТ] NPV = −9 500 000 + 1 432 000 / 1,24 + 4 036 320 / 1,24² + 5 961 575,20 / 1,24³ + 6 930 350,556 / 1,24⁴ + 10 578 478,20016 / 1,24⁵ = '+num(f['base']['npv_rub'])+' руб.')
 b.para('[РАСЧЕТ] Безубыточность после УСН: 4 200 000 / (1 − 2 640 000 / 8 800 000 − 0,06) = 6 562 500 руб. Эквивалентный набор: '+num(f['break_even_2027']['treatment_ha_passes'])+' га-проходов и '+num(f['break_even_2027']['monitoring_ha_surveys'])+' га-обследований. При изменении микса порог пересчитывается.')
 b.doc.add_heading('Приложение Б. Помесячный денежный поток',level=1);b.para('[ДОПУЩЕНИЕ] Поступления в месяц выполнения; аванс не улучшает модель. [РАСЧЕТ] Начальный остаток после CAPEX и запуска — 2 000 000 руб. Таблица включает все 12 месяцев; месячный УСН — резервирование, не официальный платежный календарь.')
 mm=f['first_year_monthly'];keys=['opening_cash_rub','revenue_collected_ex_vat_rub','variable_cost_paid_rub','fixed_opex_paid_rub','usn_cash_provision_rub','net_operating_cf_rub','closing_cash_rub']
 b.table(['Месяц','Начало','Поступления','Переменные','Постоянные','УСН','Поток','Конец'],[[r['month']]+[num(r[k]) for k in keys] for r in mm],'ДДС 2027, руб.')
 b.doc.add_heading('Приложение В. Детальная единичная экономика',level=1)
 b.para('[РАСЧЕТ] Постоянные расходы не распределяются повторно в удельные ставки. Внесение — га-проход; мониторинг — га-обследование. Препарат клиента не включен.')
 u=rows('unit_economics.csv');b.table(['Год / услуга','Тариф без НДС','Переменные','Вклад до УСН','УСН','Вклад после УСН'],[[r['year']+' / '+('Внесение' if r['service']=='treatment' else 'Мониторинг')]+[num(r[k]) for k in ['price_ex_vat_rub','variable_cost_rub','contribution_before_usn_rub','usn_per_unit_rub','contribution_after_usn_rub']] for r in u],'Единичная экономика, руб. на операционную единицу')
 b.doc.add_heading('Приложение Г. Расчет всех стресс-сценариев',level=1)
 labels={'price_minus_10pct':'Оба тарифа −10%','volume_minus_25pct':'Объемы −25%','variable_plus_20pct':'Переменные затраты +20%','one_year_delay':'Задержка запуска на год','treatment_price_900':'Внесение 900 руб./га в 2027 с индексацией'}
 for key,title in labels.items():
  b.doc.add_heading(title,level=2);b.para('[РАСЧЕТ] Самостоятельный стресс, не совмещенный с другими. t0 остается −9 500 000 руб.; горизонт не удлиняется.')
  rr=f['yearly_scenarios'][key][1:];b.table(['Год','Выручка','EBITDA','УСН','FCF','PV'],[[str(r['year'])]+[num(r[k]) for k in ['revenue_rub','ebitda_rub','usn_tax_rub','free_cash_flow_rub','present_value_rub']] for r in rr],title+' — рубли')
  m=next(x for x in f['sensitivity'] if x['scenario']==key);b.para('[РАСЧЕТ] NPV: '+num(m['npv_rub'])+' руб. Дисконтированная окупаемость: '+('не достигнута за пять лет.' if m['discounted_payback_years'] is None else num(m['discounted_payback_years'])+' года (условная интерполяция).'))
 b.doc.add_heading('Приложение Д. Реестр параметров и допущений',level=1)
 b.para('[ДОПУЩЕНИЕ] Полный экспорт параметров модели. Идентификаторы сохранены для сверки с исходными CSV; описание и единицы поясняют смысл. Пороговые налоговые значения основаны на источниках основного раздела; будущая применимость режима требует проверки.')
 aa=rows('assumptions.csv');b.table(['Идентификатор','Значение','Единица','Обоснование / ограничение'],[[r['parameter'],num(r['value']),r['unit'],r['note']] for r in aa],'Полный реестр параметров модели')
 b.doc.add_heading('Приложение Е. Инвестиционная проверка и открытые вопросы',level=1)
 review=(ROOT/'docs/FINAL_REVIEW.md').read_text();weak=review.split('## 5. Слабые места и условия инвестиционного решения')[1];b.markdown(weak)
 path=OUT/'AGROVECTOR_FULL_BUSINESS_PLAN.docx';b.doc.save(path)
 # Read-back: package, sections, every source table cell, finance and appendices.
 doc=Document(path);alltext='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
 headers=[p.text for p in doc.paragraphs if p.style.name=='Heading 1' and re.match(r'\d+\.',p.text)]
 assert len(headers)==22,headers
 assert all(f'{i}. '+t in headers for i,t in enumerate(titles,1))
 for table in b.source_tables:
  for row in table:
   for cell in row:assert cell in alltext,cell
 for marker in ['Приложение А.','Приложение Б.','Приложение В.','Приложение Г.','Приложение Д.','Приложение Е.','3 946 443,10','9 500 000','[ФАКТ]','[РАСЧЕТ]','[ДОПУЩЕНИЕ]']:assert marker in alltext,marker
 assert 'SpareForge' not in alltext and '```' not in alltext and 'flowchart LR' not in alltext
 assert len(doc.inline_shapes)==2
 report={'file':str(path),'sections_22_verified':len(headers),'appendices':6,'tables':len(doc.tables),'diagrams':len(doc.inline_shapes),'source_table_cells_preserved':True,'assumptions':len(aa),'monthly_rows':len(mm),'unit_economics_rows':len(u),'stress_scenarios':len(labels),'npv_rub':f['base']['npv_rub'],'paragraphs':len(doc.paragraphs),'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'visual_rendering':False,'note':'DOCX structurally verified by read-back; page count and pagination not asserted without Word/LibreOffice rendering.'}
 (OUT/'DOCX_VALIDATION.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
