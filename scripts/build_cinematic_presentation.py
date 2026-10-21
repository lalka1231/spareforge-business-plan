"""Cinematic AgroVector: embedded moving field scenes, editable evidence overlays."""
from pathlib import Path
import json,csv,copy,zipfile,hashlib
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1];ASSETS=ROOT/'presentations/cinematic_assets';OUT=ROOT/'presentations/AGROVECTOR_CINEMATIC.pptx'
P='http://schemas.openxmlformats.org/presentationml/2006/main';NS={'p':P}
def color(h):return RGBColor.from_string(h)
def text(s,x,y,w,h,t,size=24,bold=False,c='FFFFFF'):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=0
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(c);p.space_after=Pt(10)
 return sh
def panel(s,x,y,w,h,c='0A1B19',alpha=87000):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=color(c);sh.line.fill.background()
 rgb=sh._element.xpath('.//a:solidFill/a:srgbClr')[0];E.SubElement(rgb,'{http://schemas.openxmlformats.org/drawingml/2006/main}alpha',val=str(alpha));return sh
def table(s,headers,rows,x=.6,y=2.4,w=12.1,h=3.7,widths=None,size=15):
 tab=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h)).table
 if widths:
  for col,ww in zip(tab.columns,widths):col.width=Inches(ww)
 for i,row in enumerate([headers]+rows):
  for j,value in enumerate(row):
   cell=tab.cell(i,j);cell.text=str(value);cell.fill.solid();cell.fill.fore_color.rgb=color('264E3B' if i==0 else ('102620' if i%2 else '19372C'))
   cell.margin_left=Inches(.11);cell.margin_right=Inches(.09);cell.margin_top=Inches(.07)
   for p in cell.text_frame.paragraphs:p.font.name='Arial';p.font.size=Pt(size);p.font.bold=i==0;p.font.color.rgb=color('E3F4B0' if i==0 else 'FFFFFF')
 return tab

def build():
 manifest=json.loads((ROOT/'data/cinematic_content.json').read_text());base=Presentation(str(ROOT/'presentations/AGROVECTOR_22_SLIDES.pptx'));p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
 scene_keys=['hover','flight','detail','survey','hover','survey','wheat','survey','detail','wheat','flight','detail','survey','hover','survey','detail','survey','hover','detail','hover','flight','hover'];movieids=[]
 for i,item in enumerate(manifest,1):
  s=p.slides.add_slide(p.slide_layouts[6]);key=scene_keys[i-1];mv=s.shapes.add_movie(str(ASSETS/(key+'.mp4')),0,0,p.slide_width,p.slide_height,poster_frame_image=str(ASSETS/(key+'_poster.jpg')),mime_type='video/mp4');movieids.append(mv.shape_id)
  s.shapes.add_picture(str(ASSETS/'scrim.png'),0,0,width=p.slide_width,height=p.slide_height)
  # Clear editorial opening; every scene is visibly a product field environment.
  panel(s,.45,.42,.055,.35,'D6ED89',100000)
  text(s,.65,.39,11.9,.45,f'АГРОВЕКТОР  /  {i:02d}  /  {item["title"].upper()}',13,True,'D6ED89')
  headline='АгроВектор' if i==1 else item['headline']
  text(s,.65,1.35,7.1 if i not in (6,8,14,15,16) else 11.9,1.65,headline,54 if i==1 else 32,True)
  if i==1:
   text(s,.68,2.6,5.9,1.1,'Мониторинг посевов.\nТочечное внесение. Проверяемый отчет.',24)
   text(s,.68,4.18,6.0,.55,'Татарстан  /  План 2027–2031',16,False,'D6ED89')
   text(s,.68,5.1,6.2,1.12,'9,5 млн ₽ — стартовая потребность\nЗаказы, разрешения и деньги не подтверждены.',20)
  elif i==4:
   blocks=[('СЕГМЕНТЫ','КФХ, малые и средние хозяйства'),('ЦЕННОСТЬ','Задание → карта → внесение → отчет'),('КАНАЛЫ','Прямые продажи, агрономы'),('ОТНОШЕНИЯ','Планирование сезона, единый контакт'),('ДОХОДЫ','Внесение + отдельный мониторинг'),('РЕСУРСЫ','Одна бригада, три БВС'),('ДЕЙСТВИЯ','Продажи, допуски, работы, контроль'),('ПАРТНЕРЫ','Агроном, сервис, юрист — кандидаты'),('ЗАТРАТЫ','OPEX; CAPEX и резерв отдельно')]
   for j,(title,body) in enumerate(blocks):
    x=.65+(j%3)*4.08;y=3.0+(j//3)*1.05;panel(s,x,y,3.9,.94);text(s,x+.14,y+.1,3.6,.25,title,12,True,'D6ED89');text(s,x+.14,y+.38,3.6,.47,body,15)
  elif i==6:
   matrix=json.loads((ROOT/'data/competitive_matrix.json').read_text());table(s,matrix['slide_headers'],matrix['slide_rows'],y=2.4,h=3.8,widths=[1.5,3.7,2.5,4.4],size=13)
   text(s,.65,6.48,11.9,.3,'Заявления сайтов; доли не раскрыты. Цены регионов и комплекты несопоставимы. URL, даты и полная матрица — в заметках.',12,False,'D6ED89')
  elif i in (8,14,15,16):
   # Preserve the source's real table/diagram, not a decorative substitute.
   source=base.slides[i-1]
   if i in (8,16):
    tab=next(sh.table for sh in source.shapes if sh.has_table); headers=[c.text for c in tab.rows[0].cells];rows=[[c.text for c in r.cells] for r in list(tab.rows)[1:]]
    table(s,headers,rows,y=2.85,h=3.65,size=14 if i==8 else 12)
   else:
    panel(s,.6,2.8,12.1,3.62,alpha=93000)
    # Copy only the editable schedule/network artwork, excluding old chrome.
    copies=[];mapping={};nextid=max(sh.shape_id for sh in s.shapes)+1
    for sh in source.shapes:
     if 1.7<sh.top/914400<5.9:
      clone=copy.deepcopy(sh._element)
      for nv in clone.xpath('.//p:cNvPr'):
       mapping[nv.get('id')]=str(nextid);nv.set('id',str(nextid));nextid+=1
      for off in clone.xpath('./p:spPr/a:xfrm/a:off'):
       off.set('y',str(int(Inches(2.8)+(int(off.get('y'))-Inches(1.7))*.8)))
      for ext in clone.xpath('./p:spPr/a:xfrm/a:ext'):ext.set('cy',str(int(int(ext.get('cy'))*.8)))
      copies.append(clone)
    for clone in copies:
     for conn in clone.xpath('.//a:stCxn | .//a:endCxn'):
      if conn.get('id') in mapping:conn.set('id',mapping[conn.get('id')])
     s.shapes._spTree.insert_element_before(clone,'p:extLst')
    text(s,.7,6.45,11.8,.3,'План, не выполненные этапы. Готовность и разрешения — обязательные условия.',12,False,'D6ED89')
  elif i in (17,18):
   text(s,7.72,2.7,4.8,.25,'МЛН РУБ.  /  '+('ВЫРУЧКА БЕЗ НДС' if i==17 else 'NPV; СТАВКА 24%'),11,True,'D6ED89')
   panel(s,7.45,2.99,5.15,3.35,alpha=92000)
   srcchart=next(sh.chart for sh in base.slides[i-1].shapes if sh.has_chart)
   from pptx.chart.data import CategoryChartData
   from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION
   data=CategoryChartData();data.categories=[c.label for c in srcchart.plots[0].categories]
   for series in srcchart.series:data.add_series(series.name,list(series.values))
   ch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED if i==17 else XL_CHART_TYPE.BAR_CLUSTERED,Inches(7.7),Inches(3.05),Inches(4.65),Inches(2.9),data).chart
   ch.has_legend=False
   if i==17:
    for j,(label,tint) in enumerate([('Выручка','D6ED89'),('EBITDA','E5B26A'),('FCF','A7CAB9')]):
     panel(s,7.72+j*1.6,6.15,.1,.1,tint,100000);text(s,7.9+j*1.6,6.08,1.4,.3,label,11,False,'FFFFFF')
   for ax in (ch.category_axis,ch.value_axis):ax.tick_labels.font.size=Pt(11);ax.tick_labels.font.color.rgb=color('FFFFFF')
   for j,series in enumerate(ch.series):series.format.fill.solid();series.format.fill.fore_color.rgb=color(['D6ED89','E5B26A','A7CAB9'][j%3])
   text(s,.7,3.35,6.4,2.8,item['visible_text'],22)
   if i==18:text(s,.7,6.2,11.5,.55,'Ресурсный стресс: 20 полевых дней/месяц → NPV 2,310 млн ₽. Это допущение, не подтвержденная мощность.',15,False,'D6ED89')
  else:
   lines=item['visible_text'].split('\n')
   for j,line in enumerate(lines):
    y=3.6+j*.88;panel(s,.65,y,6.7,.77,alpha=78000);text(s,.82,y+.12,6.35,.62,line,19 if len(line)>90 else 22)
  # Footer always names the illustration and separates it from project evidence.
  panel(s,0,7.04,13.34,.46,alpha=88000)
  text(s,.65,7.13,11.2,.22,'Визуализация концепции + сток Pexels. Не съемка AgroVector.  /  Редакция 20.10.2026',10,False,'CCDACF')
  text(s,12.05,7.11,.7,.23,f'{i:02d}/22',11,True,'D6ED89')
  s.notes_slide.notes_text_frame.text=base.slides[i-1].notes_slide.notes_text_frame.text+'\n\nКИНЕМАТОГРАФИЧЕСКАЯ СЦЕНА: '+key+'. Иллюстративный сток Pexels и собственная 3D-визуализация; не свидетельство реальных работ AgroVector. Полная матрица и доказательства сохранены в исходных заметках. Права/URL: presentations/cinematic_assets/asset_ledger.json.\n'+item['scene']
 p.save(str(OUT))
 # Final timing is patched AFTER all movie shapes, never through build_motion.py.
 with zipfile.ZipFile(OUT) as z:parts={n:z.read(n) for n in z.namelist()}
 for i,spid in enumerate(movieids,1):
  n=f'ppt/slides/slide{i}.xml';tree=E.fromstring(parts[n]);tim=tree.find('p:timing',NS)
  assert tim is not None
  for media in tim.findall('.//p:cMediaNode',NS):
   media.set('vol','0');ct=media.find('p:cTn',NS);ct.set('repeatCount','indefinite');ct.set('fill','hold');ct.find('p:stCondLst/p:cond',NS).set('delay','0')
  tr=E.Element('{'+P+'}transition',spd='slow',advClick='1');E.SubElement(tr,'{'+P+'}push',dir='u');tree.insert(list(tree).index(tim),tr)
  parts[n]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
  for n,b in parts.items():z.writestr(n,b)
 check=Presentation(str(OUT));assert len(check.slides)==22
 with zipfile.ZipFile(OUT) as z:
  assert z.testzip() is None
  movies=[n for n in z.namelist() if n.endswith('.mp4')];assert len(movies)==5
  for i in range(1,23):
   t=E.fromstring(z.read(f'ppt/slides/slide{i}.xml'));assert t.find('p:transition/p:push',NS) is not None
   assert t.find('.//p:cMediaNode/p:cTn/p:stCondLst/p:cond',NS).get('delay')=='0'
 report={'file':str(OUT),'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'slides':22,'embedded_video_files':len(movies),'video_scenes':22,'vertical_push_transitions':22,'autoplay_loop_xml':True,'native_slideshow_playback':'not yet verified','scene_keys':scene_keys,'limitations':['Video scenes are edits/composites, not one uninterrupted camera take','Pexels images do not demonstrate AgroVector operations','Synthetic drone is an illustrative non-manufacturer-specific model','No claim that sprayed stock liquid is fertilizer','Physical hovering and passing are encoded in MP4; no Morph required to animate the drone']}
 (ROOT/'presentations/CINEMATIC_VALIDATION.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':build()
