"""Independent read-back of final cinematic package and actual video decodes."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
import json,hashlib,tempfile,subprocess,imageio_ffmpeg,pymupdf
ROOT=Path(__file__).resolve().parents[1];PATH=ROOT/'presentations/AGROVECTOR_CINEMATIC.pptx';NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
def verify():
 p=Presentation(str(PATH));base=Presentation(str(ROOT/'presentations/AGROVECTOR_22_SLIDES.pptx'));assert len(p.slides)==22
 for i,(a,b) in enumerate(zip(base.slides,p.slides),1):
  assert a.notes_slide.notes_text_frame.text in b.notes_slide.notes_text_frame.text,(i,'source notes lost')
  assert 'Не съемка AgroVector' in '\n'.join(s.text for s in b.shapes if s.has_text_frame)
 matrix=json.loads((ROOT/'data/competitive_matrix.json').read_text());t=next(s.table for s in p.slides[5].shapes if s.has_table);assert [[c.text for c in r.cells] for r in list(t.rows)[1:]]==matrix['slide_rows']
 for i in (8,16):
  ta=next(s.table for s in base.slides[i-1].shapes if s.has_table);tb=next(s.table for s in p.slides[i-1].shapes if s.has_table);assert [[c.text for c in r.cells] for r in ta.rows]==[[c.text for c in r.cells] for r in tb.rows]
 for i in (17,18):
  ca=next(s.chart for s in base.slides[i-1].shapes if s.has_chart);cb=next(s.chart for s in p.slides[i-1].shapes if s.has_chart)
  for sa,sb in zip(ca.series,cb.series):assert sa.name==sb.name and all(abs(a-b)<1e-10 for a,b in zip(sa.values,sb.values))
 ff=imageio_ffmpeg.get_ffmpeg_exe();decodes=[]
 with ZipFile(PATH) as z:
  assert z.testzip() is None;movies=[n for n in z.namelist() if n.endswith('.mp4')];assert len(movies)==5
  for i in range(1,23):
   tree=E.fromstring(z.read(f'ppt/slides/slide{i}.xml'));ids=[x.get('id') for x in tree.findall('.//p:cNvPr',NS)];assert len(ids)==len(set(ids));assert tree.find('p:transition/p:push',NS).get('dir')=='u'
   ct=tree.find('.//p:video/p:cMediaNode/p:cTn',NS);assert ct.get('repeatCount')=='indefinite'
   main=tree.find('.//p:cTn[@nodeType="mainSeq"]',NS)
   if main is None:assert ct.find('p:stCondLst/p:cond',NS).get('delay')=='0'
   else:
    assert main.find('.//p:cmd[@cmd="playFrom(0.0)"]',NS) is not None
    assert main.find('.//p:cTn[@nodeType="withEffect"]',NS) is not None
    assert all(c.get('delay')!='indefinite' for c in main.findall('.//p:cond',NS))
   assert tree.find('.//p:video/p:cMediaNode/p:tgtEl/p:spTgt',NS).get('spid') in ids
  with tempfile.TemporaryDirectory() as tmp:
   for name in movies:
    dest=Path(tmp)/Path(name).name;dest.write_bytes(z.read(name));result=subprocess.run([ff,'-v','error','-i',str(dest),'-f','null','-'],stdout=subprocess.PIPE,stderr=subprocess.PIPE);assert result.returncode==0,(name,result.stderr.decode());decodes.append({'embedded':name,'bytes':dest.stat().st_size,'full_decode':True})
 pdf=ROOT/'review/cinematic/AGROVECTOR_CINEMATIC.pdf';assert len(pymupdf.open(pdf))==22
 report={'verified':True,'sha256':hashlib.sha256(PATH.read_bytes()).hexdigest(),'slides':22,'video_scenes':22,'unique_embedded_videos':5,'push_transitions':22,'immediate_looping_media':22,'source_notes_preserved':22,'tables_and_charts_consistent':True,'decoded_media':decodes,'pdf_pages':22,'native_playback_limit':'Native autoplay/loop settings true on all 22 slides after native save; With Previous commands persist. Pixel-level playback is not captured.'}
 (ROOT/'presentations/CINEMATIC_ACCEPTANCE.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':verify()
