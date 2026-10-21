"""Read back the exact teacher-reference PPTX and independently verify model links."""
from pathlib import Path
from zipfile import ZipFile
import json,csv,hashlib,math,re
from pptx import Presentation
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1]
def main():
 path=ROOT/'presentations/AGROVECTOR_PHOTO_MORPH.pptx';p=Presentation(str(path));assert len(p.slides)==22
 def text(n):
  return '\n'.join(sh.text if sh.has_text_frame else '\n'.join(c.text for r in sh.table.rows for c in r.cells) if sh.has_table else '' for sh in p.slides[n-1].shapes)
 assert 'LEAN CANVAS' in text(4)
 for n in ['1. Сегменты','2. Проблемы','3. Уникальная','4. Решение','5. Каналы','6. Потоки','7. Структура','8. Ключевые','9. Преимущество']:assert n in text(4),n
 for key in ['SO:','ST:','WO:','WT:']:assert key in text(7)
 assert len([s for s in p.slides[7].shapes if s.has_table][0].table.columns)==4
 assert 'МЕНДЕЛОУ' in text(10)
 for t in ['Высокая власть / низкий','Высокая власть / высокий','Низкая власть / низкий','Низкая власть / высокий']:assert t in text(10)
 tasks=list(csv.DictReader((ROOT/'data/project_plan.csv').open()))
 wbs=text(13)
 for r in tasks:assert r['id']+'  ' in wbs,r['id']
 tab=next(sh.table for sh in p.slides[15].shapes if sh.has_table);assert len(tab.rows)==14
 for r in list(tab.rows)[1:]:
  assert sum('A' in c.text.split('/') for c in list(r.cells)[1:])==1
  assert sum('R' in c.text.split('/') for c in list(r.cells)[1:])>=1
 for task in ['I','K']:
  source=next(r for r in tasks if r['id']==task);assert source['R']=='Второй оператор' and source['A']=='Операционный руководитель'
  r=next(r for r in list(tab.rows)[1:] if r.cells[0].text.startswith(task+' '));assert r.cells[6].text=='A' and r.cells[7].text=='R'
 market=json.loads((ROOT/'data/market_revision.json').read_text());supp=json.loads((ROOT/'data/supplemental_summary.json').read_text())
 assert market['tam_rub']==79298000*.15*(1.5*1300+.5*250)==supp['market']['tam_rub']
 assert 'Оператор 2' in text(16)
 fin=json.loads((ROOT/'data/finance_summary.json').read_text());annual=fin['yearly_scenarios']['base'];npv=sum(r['free_cash_flow_rub']/1.24**r['period'] for r in annual);assert math.isclose(npv,fin['base']['npv_rub'],abs_tol=.01)
 payroll=list(csv.DictReader((ROOT/'data/payroll2027.csv').open()))
 # Independently evaluate the explicitly chosen 3 FTE staffing scenario.
 employer=3*100000*12*(1+.30+.002)
 assert math.isclose(annual[1]['fixed_opex_rub'],employer+1200000,abs_tol=.01)
 assert math.isclose(annual[1]['ebitda_rub'],8800000-2640000-employer-1200000,abs_tol=.01)
 for i in range(1,23):
  sec=json.loads((ROOT/'data/sections.json').read_text())[str(i)]
  note=p.slides[i-1].notes_slide.notes_text_frame.text
  # Document normalizes Russian word/number spacing; ignore whitespace for preservation.
  assert re.sub(r'\s+','',sec) in re.sub(r'\s+','',note),f'notes {i}'
 with ZipFile(path) as z:
  assert z.testzip() is None;assert not any(n.lower().endswith('.mp4') for n in z.namelist())
  ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','m':'http://schemas.microsoft.com/office/powerpoint/2015/09/main'}
  for i in range(1,23):
   t=E.fromstring(z.read(f'ppt/slides/slide{i}.xml'));assert t.xpath('.//m:morph',namespaces=ns)
   assert not t.xpath('//*[local-name()="videoFile" or local-name()="video"]')
   if i==2:
    bg=next(s.shape_id for s in p.slides[1].shapes if s.name=='!!AV_background')
    assert not t.xpath(f'.//p:animMotion//p:spTgt[@spid="{bg}"]',namespaces=ns)
 report={'verified':True,'slides':22,'lean_canvas_blocks':9,'tows_strategies':4,'porter_columns':4,'mendelow_quadrants':4,'wbs_packages':len(tasks),'raci_rows':len(tab.rows)-1,'payroll_employer_rub':employer,'npv_independent_rub':npv,'ebitda_2027_rub':annual[1]['ebitda_rub'],'notes_current':22,'morph':22,'video_files':0,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'visual_playback':'not certified by structural read-back'}
 (ROOT/'presentations/REFERENCE_REVISION_ACCEPTANCE.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
